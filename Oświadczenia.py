import os
import re
import json
import pandas as pd
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
import win32com.client as win32
from collections import defaultdict
import tkinter as tk
from tkinter import ttk, messagebox
from tkinter.filedialog import asksaveasfilename
from datetime import datetime
import threading
import pythoncom
import numpy as np
import ctypes
from ctypes import wintypes, windll
import io
from reportlab.pdfgen import canvas
from PIL import Image, ImageFilter, ImageEnhance
from PyPDF2 import PdfReader, PdfWriter
from pdf2image import convert_from_path
import tempfile
import random
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
import sharepoint_sync

# ------------------------- Get Desktop Path -------------------------
CSIDL_DESKTOP = 0
_SHGetFolderPath = windll.shell32.SHGetFolderPathW
_SHGetFolderPath.argtypes = [wintypes.HWND,
                             ctypes.c_int,
                             wintypes.HANDLE,
                             wintypes.DWORD, wintypes.LPCWSTR]
path_buf = ctypes.create_unicode_buffer(wintypes.MAX_PATH)
result = _SHGetFolderPath(0, CSIDL_DESKTOP, 0, 0, path_buf)
desktop_path = path_buf.value
print("Pulpit:", desktop_path)

APP_FOLDER_NAME = "Oświadczenia aplikacja"

def find_app_folder():
    candidates = []

    for env in ("OneDriveCommercial", "OneDrive", "OneDriveConsumer"):
        base = os.environ.get(env)
        if base:
            candidates.append(os.path.join(base, APP_FOLDER_NAME))

    # 2) Pulpit (obejmuje też Pulpit przekierowany do OneDrive)
    candidates.append(os.path.join(desktop_path, APP_FOLDER_NAME))

    # 3) Katalog profilu użytkownika
    userprofile = os.environ.get("USERPROFILE")
    if userprofile:
        candidates.append(os.path.join(userprofile, APP_FOLDER_NAME))

    # Zwróć pierwszą lokalizację, która realnie istnieje
    for path in candidates:
        if os.path.isdir(path):
            print("Znaleziono folder aplikacji:", path)
            return path

    # Nic nie znaleziono – zwróć najlepsze przypuszczenie (pierwsze z listy)
    fallback = candidates[0] if candidates else os.path.join(desktop_path, APP_FOLDER_NAME)
    print("UWAGA: nie znaleziono folderu aplikacji. Używam:", fallback)
    return fallback

base_dir = find_app_folder()

# Signatures folder (client data comes from SharePoint, see sharepoint_sync.py)
signature_dir = os.path.join(base_dir, "podpisy")
# Last used choices, remembered per user
settings_path = os.path.join(sharepoint_sync.DATA_DIR, "ustawienia.json")

# ------------------------- Background Work -------------------------
def run_in_background(work, on_done):
    """Run work() in a worker thread, then call on_done(result, error) in the main thread."""
    outcome = {}

    def worker():
        try:
            outcome['result'] = work()
        except Exception as e:
            outcome['error'] = e

    thread = threading.Thread(target=worker, daemon=True)
    thread.start()

    def wait():
        # tkinter may only be used from the main thread, so wait for the worker here
        # instead of updating the window from it.
        if thread.is_alive():
            root.after(100, wait)
        else:
            on_done(outcome.get('result'), outcome.get('error'))
    root.after(100, wait)

def show_busy(text):
    """Small window with a moving progress bar that blocks the app until it is destroyed."""
    busy = tk.Toplevel(root)
    busy.title("Proszę czekać")
    busy.resizable(False, False)
    busy.transient(root)
    busy.protocol("WM_DELETE_WINDOW", lambda: None)
    ttk.Label(busy, text=text).pack(padx=20, pady=(15, 5))
    progress = ttk.Progressbar(busy, mode='indeterminate', length=250)
    progress.pack(padx=20, pady=(0, 15))
    progress.start(10)
    busy.grab_set()
    return busy

# ------------------------- Settings -------------------------
def load_settings():
    try:
        with open(settings_path, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}

def save_settings(**settings):
    try:
        os.makedirs(os.path.dirname(settings_path), exist_ok=True)
        with open(settings_path, "w", encoding="utf-8") as f:
            json.dump(settings, f, ensure_ascii=False)
    except OSError:
        pass  # remembering the choices is only a convenience

# ------------------------- SharePoint Data -------------------------
def load_data():
    """Load the local copy of the SharePoint list (kept in the user's AppData)."""
    global data, client_names, data_synced_at
    try:
        data, data_synced_at = sharepoint_sync.load()
    except Exception as e:
        messagebox.showerror("Error", f"Błąd wczytywania danych: {e}")
        data, data_synced_at = pd.DataFrame(columns=list(sharepoint_sync.COLUMNS)), None
    # The same client is sometimes typed with different spacing or letter case: group
    # those rows under one key and list the client once, under its most common spelling.
    data['client_key'] = data['Nazwa firmy'].fillna('').map(normalize_name)
    spellings = (data.groupby(['client_key', 'Nazwa firmy']).size()
                 .sort_values(ascending=False, kind='stable').reset_index())
    client_names = sorted((" ".join(name.split()) for name in
                           spellings.drop_duplicates('client_key')['Nazwa firmy']), key=str.casefold)
    client_name_combobox['values'] = client_names

def client_rows(client_entries):
    """All rows of the given client entries, whichever spelling of the name they use."""
    return data[data['client_key'].isin({normalize_name(name) for name in client_entries})]

def show_data_status(note=""):
    if data_synced_at:
        text = "Dane z SharePoint z " + datetime.fromisoformat(data_synced_at).strftime("%d.%m.%Y %H:%M")
    else:
        text = "Brak danych z SharePoint"
    status_var.set(text + note)

def refresh_data(full=True):
    """Update the local copy from SharePoint in the background.

    full=True ("Odśwież dane" button) downloads the whole list, otherwise only the changes."""
    refresh_button.config(state=tk.DISABLED)
    status_var.set("Pobieranie danych z SharePoint..." if full else "Sprawdzanie zmian w SharePoint...")
    parent_window = int(root.wm_frame(), 16)  # the sign-in window, if needed, opens over the app

    def on_done(result, error):
        refresh_button.config(state=tk.NORMAL)
        if error:
            show_data_status(" (nie udało się zaktualizować)")
            if full or data.empty:
                messagebox.showerror("Error", f"Nie udało się pobrać danych z SharePoint: {error}")
            return
        print("Synchronizacja z SharePoint:", result)
        load_data()
        show_data_status()

    run_in_background(lambda: sharepoint_sync.sync(full=full, parent_window_handle=parent_window), on_done)

# ------------------------- Helper Functions -------------------------
def add_signature(sheet, image_path, cell):
    img = OpenpyxlImage(image_path)
    sheet.add_image(img, cell)

def normalize_name(name):
    # Names are compared ignoring letter case and extra spaces: SharePoint, the signature
    # files and the client names typed in SharePoint do not always agree on them
    # (e.g. "Michał Augustyn" vs "Michał  Augustyn_czarny_1.png").
    return " ".join(name.split()).casefold()

def find_signatures(signer, color):
    if not os.path.isdir(signature_dir):
        return []
    wanted = {normalize_name(f"{signer}_{color}_{j}.png") for j in range(1, 4)}
    return [os.path.join(signature_dir, file_name) for file_name in os.listdir(signature_dir)
            if normalize_name(file_name) in wanted]

def safe_filename(name):
    """The name with the characters Windows does not allow in file and folder names replaced."""
    return re.sub(r'[<>:/\\|?*\x00-\x1f]', '_', name.replace('"', "'")).strip()

def shorten_for_path(name, limit=60):
    """Client name cut to at most `limit` characters for folder and file names, so that the
    full path stays under Windows' 260-character limit. The statement keeps the full name."""
    if len(name) <= limit:
        return name
    cut = name[:limit]
    if name[limit] != " " and " " in cut:
        cut = cut.rsplit(" ", 1)[0]  # do not cut a word in half
    return cut.rstrip(" ,.-")

def statement_scope(audit_type):
    """'jednostkowe', 'skonsolidowane' or both, depending on the selected statement type."""
    text = audit_type.lower()
    scopes = [scope for stem, scope in (("jednostkow", "jednostkowe"), ("skonsolidowan", "skonsolidowane"))
              if stem in text]
    return " i ".join(scopes)

def parse_date(text):
    """Date from DD.MM.YYYY text, or None when the text is not such a date."""
    try:
        return datetime.strptime(text.strip(), '%d.%m.%Y').date()
    except ValueError:
        return None

# ------------------------- Scanned Effect Function -------------------------
def add_scanned_effect(img):
    # Slight random rotation
    angle = np.random.uniform(-0.5, 0.4)
    img = img.rotate(angle, expand=1, fillcolor=(255,255,255))
    # Add slight blur
    img = img.filter(ImageFilter.GaussianBlur(radius=0.7))
    # Add noise
    arr = np.asarray(img, dtype=np.int16)
    arr += np.random.default_rng().normal(0, 8, arr.shape).astype(np.int16)
    img = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))
    # Adjust contrast and brightness
    enhancer = ImageEnhance.Contrast(img)
    img = enhancer.enhance(1.15)
    enhancer = ImageEnhance.Brightness(img)
    img = enhancer.enhance(1.05)
    return img

# ------------------------- Client Record Helpers -------------------------
def normalize_task_type(value):
    return str(value).strip().lower() if pd.notnull(value) else ''

def get_client_records(client_entries):
    records = []
    for _, row in client_rows(client_entries).iterrows():
        records.append({
            'signer_name': row['Osoba odpowiedzialna'],
            'task_type': row['Typ zadania'],
            'appearance_date': row['Data rozpoczęcia'],
            'type_audit': row['Rodzaj sprawozdania'],
        })
    return records

def compute_earliest_date(signer_data_list, client_name):
    dates = [entry['appearance_date'] for entry in signer_data_list
             if normalize_task_type(entry['task_type']) != "oświadczenie"]
    if not dates:
        messagebox.showerror("Error", f"Nie znaleziono odpowiednich danych dla: {client_name} (zadania poza 'Oświadczenie')")
        return None
    return min(dates)

# ------------------------- Signer Selection Dialog -------------------------
def open_signer_selection_dialog(client_name, client_entries, dzien_otw_bil, dzien_bil, audit_type,
                                 data_podpisu_umowy, data_podpisu_badania):
    """client_entries: the client's entries in the list - just client_name, or several consolidated ones."""
    client_records = client_rows(client_entries).copy()
    if client_records.empty:
        messagebox.showerror("Error", f"Nie znaleziono danych dla klienta: {client_name}")
        return

    signer_data_list = get_client_records(client_entries)

    computed_earliest_date = compute_earliest_date(signer_data_list, client_name)
    if computed_earliest_date is None:
        return

    # Only tasks other than "Oświadczenie" make someone a signer, both here and when a row is selected.
    signer_tasks = defaultdict(list)
    for entry in signer_data_list:
        if normalize_task_type(entry['task_type']) != "oświadczenie" and isinstance(entry['signer_name'], str):
            signer_tasks[entry['signer_name']].append({
                'task_type': entry['task_type'],
                'appearance_date': entry['appearance_date']
            })
    valid_signers = list(signer_tasks)

    # People without a signature file in the chosen colour are shown, but cannot be selected.
    signature_color = signature_color_var.get()
    missing_signature = {signer for signer in valid_signers if not find_signatures(signer, signature_color)}

    sel_dialog = tk.Toplevel(root)
    sel_dialog.title("Wybór daty i podpisujących")
    sel_dialog.geometry("800x800")

    # Table frame
    table_frame = ttk.LabelFrame(sel_dialog, text="Rekordy klienta", padding="10")
    table_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

    # Create Treeview with scrollbars
    tree_frame = ttk.Frame(table_frame)
    tree_frame.pack(fill=tk.BOTH, expand=True)
    vsb = ttk.Scrollbar(tree_frame, orient="vertical")
    vsb.pack(side="right", fill="y")
    hsb = ttk.Scrollbar(tree_frame, orient="horizontal")
    hsb.pack(side="bottom", fill="x")
    columns = ("Data rozpoczęcia", "Osoba odpowiedzialna", "Typ zadania", "Rodzaj sprawozdania")
    consolidated = len(client_entries) > 1
    if consolidated:
        columns += ("Nazwa firmy",)  # which of the consolidated entries a record comes from
    tree = ttk.Treeview(tree_frame, columns=columns, show="headings",
                        yscrollcommand=vsb.set, xscrollcommand=hsb.set)
    for col in columns:
        tree.heading(col, text=col)
        tree.column(col, width=150)
    tree.pack(fill=tk.BOTH, expand=True)
    vsb.config(command=tree.yview)
    hsb.config(command=tree.xview)

    client_records['Data rozpoczęcia'] = pd.to_datetime(client_records['Data rozpoczęcia'], errors='coerce')
    client_records = client_records.sort_values(by='Data rozpoczęcia')

    # Insert records, remembering each row's date (the shown DD.MM.YYYY text is not parsed back)
    row_dates = {}
    for idx, row in client_records.iterrows():
        start = row['Data rozpoczęcia']
        date_str = start.strftime("%d.%m.%Y") if pd.notnull(start) else ''
        values = [date_str, row['Osoba odpowiedzialna'], row['Typ zadania'], row['Rodzaj sprawozdania']]
        if consolidated:
            values.append(row['Nazwa firmy'])
        item = tree.insert("", "end", values=values)
        row_dates[item] = start

    # Options frame
    options_frame = ttk.LabelFrame(sel_dialog, text="Wybór osób podpisujących i daty", padding="10")
    options_frame.pack(fill=tk.X, padx=10, pady=10)
    ttk.Label(options_frame, text="Osoby podpisujące:").grid(row=0, column=0, sticky=tk.W)
    checkbox_frame = ttk.Frame(options_frame)
    checkbox_frame.grid(row=1, column=0, sticky=tk.W, padx=5, pady=5)
    signer_vars = {}
    for i, signer in enumerate(valid_signers):
        has_signature = signer not in missing_signature
        var = tk.IntVar(value=int(has_signature))
        signer_vars[signer] = var
        cb = ttk.Checkbutton(checkbox_frame, variable=var,
                             text=signer if has_signature else f"{signer} (brak podpisu: {signature_color})",
                             state=tk.NORMAL if has_signature else tk.DISABLED)
        cb.grid(row=i // 3, column=i % 3, sticky=tk.W, padx=5, pady=2)

    ttk.Label(options_frame, text="Data kontaktu (pierwszego zlecenia):").grid(row=2, column=0, sticky=tk.W, padx=5, pady=5)
    custom_date_var = tk.StringVar(value=computed_earliest_date.strftime("%d.%m.%Y"))
    if na_dzien_podpisu_var.get() == 1:
        custom_date_entry = ttk.Entry(options_frame, textvariable=custom_date_var, width=15, state='disabled')
    else:
        custom_date_entry = ttk.Entry(options_frame, textvariable=custom_date_var, width=15)
    custom_date_entry.grid(row=2, column=1, sticky=tk.W, padx=5, pady=5)

    # New frame for client name override
    override_frame = ttk.Frame(options_frame)
    override_frame.grid(row=3, column=0, columnspan=2, sticky=tk.W, padx=5, pady=5)
    ttk.Label(override_frame, text="Zmień nazwę klienta (opcjonalnie):").grid(row=0, column=0, sticky=tk.W)
    client_name_override_var = tk.StringVar(value=client_name)
    ttk.Entry(override_frame, textvariable=client_name_override_var, width=30).grid(row=0, column=1, sticky=tk.W, padx=5)

    def on_tree_select(event):
        item = tree.focus()
        if not item:
            return
        start = row_dates[item]
        if pd.isnull(start):
            custom_date_var.set('')
            return
        custom_date_var.set(start.strftime("%d.%m.%Y"))
        # Check the signers having a task (other than "Oświadczenie") on or after the selected day
        # and uncheck the others
        selected_date = start.normalize()
        for signer, var in signer_vars.items():
            if signer not in missing_signature:
                var.set(int(any(task['appearance_date'] >= selected_date for task in signer_tasks[signer])))
    tree.bind("<<TreeviewSelect>>", on_tree_select)

    # OK and Cancel buttons
    def on_ok():
        selected_signers = [signer for signer, var in signer_vars.items() if var.get() == 1]
        if not selected_signers:
            messagebox.showerror("Error", "Musisz wybrać co najmniej jedną osobę podpisującą.", parent=sel_dialog)
            return
        contact_text = custom_date_var.get().strip()
        contact_date = parse_date(contact_text) if contact_text else computed_earliest_date.date()
        if contact_date is None:
            messagebox.showerror("Error", "Nieprawidłowy format daty kontaktu. Użyj formatu DD.MM.YYYY.",
                                 parent=sel_dialog)
            return
        sel_dialog.destroy()
        # Read the override client name for display only.
        display_client_name = client_name_override_var.get().strip()
        # Pass the original client name (for filtering) and the override (for display)
        process_form(client_name, display_client_name, dzien_otw_bil, dzien_bil, audit_type,
                     data_podpisu_umowy, data_podpisu_badania, selected_signers, signature_color, contact_date)
    def on_cancel():
        sel_dialog.destroy()
    btn_frame = ttk.Frame(sel_dialog)
    btn_frame.pack(pady=10)
    ttk.Button(btn_frame, text="OK", command=on_ok).grid(row=0, column=0, padx=10)
    ttk.Button(btn_frame, text="Anuluj", command=on_cancel).grid(row=0, column=1, padx=10)

# ------------------------- Process Form -------------------------
def process_form(selected_client, display_client, dzien_otw_bil, dzien_bil, audit_type,
                 data_podpisu_umowy, data_podpisu_badania, selected_signers, signature_color, contact_date):
    skip_second_signature = na_dzien_podpisu_var.get() == 1
    if skip_second_signature:
        template_file = "Szablon dzien podpisu.xlsx"
        name_cell = 'M6'
        date_cell = 'N9'
        dzien_otw_bil_cell = 'N7'
        dzien_bil_cell = 'N8'
        audit_type_cell = 'N10'
        data_podpisu_umowy_cell = 'N12'
        data_podpisu_badania_cell = 'N13'
    else:
        template_file = "Szablon.xlsx"
        name_cell = 'L2'
        date_cell = 'M5'
        dzien_otw_bil_cell = 'M3'
        dzien_bil_cell = 'M4'
        audit_type_cell = 'M6'
        data_podpisu_umowy_cell = 'M8'

    signer_name_col = 'D'
    signature1_col, signature1_row = 'F', 22
    signature2_col, signature2_row = 'F', 35

    template_path = os.path.join(base_dir, template_file)
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    # Write the client name to Excel – use display_client if provided, otherwise the original
    final_client_name = display_client if display_client else selected_client
    if grupa_kapitalowa_var.get():
        ws[name_cell] = "Grupa kapitałowa " + final_client_name
    else:
        ws[name_cell] = final_client_name

    ws[date_cell] = contact_date
    ws[dzien_otw_bil_cell] = dzien_otw_bil
    ws[dzien_bil_cell] = dzien_bil
    ws[audit_type_cell] = audit_type
    ws[data_podpisu_umowy_cell] = data_podpisu_umowy
    if skip_second_signature:
        ws[data_podpisu_badania_cell] = data_podpisu_badania

    for i, signer in enumerate(selected_signers):
        available_signatures = find_signatures(signer, signature_color)
        if not available_signatures:
            print(f"Podpis dla {signer} nie został znaleziony.")
            continue

        signature_path = random.choice(available_signatures)
        ws[f'{signer_name_col}{signature1_row + i}'] = signer
        add_signature(ws, signature_path, f'{signature1_col}{signature1_row + i}')
        if not skip_second_signature:
            ws[f'{signer_name_col}{signature2_row + i}'] = signer
            add_signature(ws, signature_path, f'{signature2_col}{signature2_row + i}')

    for row in ws.iter_rows(min_row=22, max_row=42, min_col=4, max_col=4):
        if row[0].row == 32:
            continue
        if row[0].value is None:
            ws.row_dimensions[row[0].row].hidden = True

    # Print area of "Szablon dzien podpisu.xlsx" / "Szablon.xlsx"
    print_area = 'A1:J32' if skip_second_signature else 'A1:I31'
    ws.print_area = print_area

    if grupa_kapitalowa_var.get():
        output_client_name = "Grupa kapitałowa " + shorten_for_path(selected_client)
    else:
        output_client_name = shorten_for_path(selected_client)

    # e.g. "Oświadczenie_Firma S.A._jednostkowe_na dzień SzB.pdf"
    name_parts = [f'Oświadczenie_{output_client_name}', statement_scope(audit_type)]
    if skip_second_signature:
        name_parts.append('na dzień SzB')
    initial_name = '_'.join(part for part in name_parts if part) + '.pdf'

    # Create the folder only once if it does not exist.
    folder_path = os.path.join(desktop_path, safe_filename(f"Oświadczenia_{shorten_for_path(selected_client)}"))
    os.makedirs(folder_path, exist_ok=True)

    # Set the default directory for the save dialog to folder_path.
    output_pdf_path = asksaveasfilename(
        defaultextension=".pdf",
        filetypes=[("PDF files", "*.pdf")],
        initialdir=folder_path,
        initialfile=safe_filename(initial_name)
    )
    if not output_pdf_path:
        messagebox.showinfo("Anulowano", "Zapis PDF został anulowany.")
        return
    output_pdf_path = os.path.normpath(output_pdf_path)

    busy = show_busy("Generowanie oświadczenia, proszę czekać...")

    def on_done(result, error):
        busy.destroy()
        if error:
            messagebox.showerror("Error", f"Nie udało się wyeksportować PDF: {error}")
        elif messagebox.askyesno("Sukces", f"Oświadczenie dla {selected_client} zostało wypełnione i zapisane."
                                           "\n\nOtworzyć plik?"):
            os.startfile(output_pdf_path)

    run_in_background(lambda: export_pdf(wb, print_area, output_pdf_path), on_done)

# ------------------------- PDF Export and Flattening -------------------------
A4_WIDTH_PX = 2480  # A4 width at 300 dpi

def export_pdf(wb, print_area, output_pdf_path):
    """Save the filled workbook as a PDF that looks scanned. Runs in a worker thread."""
    # Working files go to a temporary folder, not to the folder the app was started from.
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
        excel_path = os.path.join(temp_dir, "oswiadczenie.xlsx")
        pdf_path = os.path.join(temp_dir, "oswiadczenie.pdf")
        wb.save(excel_path)
        excel_to_pdf(excel_path, pdf_path, print_area)
        flatten_pdf(pdf_path, output_pdf_path)

def excel_to_pdf(excel_path, pdf_path, print_area):
    pythoncom.CoInitialize()
    excel = win32.Dispatch("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    wb_pdf = None
    try:
        wb_pdf = excel.Workbooks.Open(excel_path)
        ws_pdf = wb_pdf.Worksheets(1)
        # Set print area using Excel's COM interface
        ws_pdf.PageSetup.PrintArea = print_area
        ws_pdf.ExportAsFixedFormat(0, pdf_path)
    finally:
        if wb_pdf is not None:
            wb_pdf.Close(False)
        excel.Application.Quit()
        pythoncom.CoUninitialize()

def flatten_pdf(input_pdf_path, output_pdf_path):
    # Convert each page into an image at A4 size (300 dpi) and rebuild the PDF from the images.
    images = convert_from_path(input_pdf_path, size=(A4_WIDTH_PX, None))
    a4_width, a4_height = A4  # A4 page dimensions in points (approx. 595x842)
    c = canvas.Canvas(output_pdf_path, pagesize=A4)
    for img in images:
        img = add_scanned_effect(img)
        width, height = img.size  # in pixels

        # Calculate scale to fit image into A4 while preserving aspect ratio.
        scale = min(a4_width / width, a4_height / height)
        scaled_width = width * scale
        scaled_height = height * scale

        # Center the image on the A4 page.
        x = (a4_width - scaled_width) / 2
        y = (a4_height - scaled_height) / 2

        # Embed the page as JPEG, like a scanner does: far smaller and faster than PNG.
        page = io.BytesIO()
        img.save(page, 'JPEG', quality=85)
        page.seek(0)
        c.drawImage(ImageReader(page), x, y, width=scaled_width, height=scaled_height)
        c.showPage()
    c.save()

# ------------------------- Main Form Functions -------------------------
def submit_form():
    client_name = client_name_var.get().strip()
    year = year_var.get().strip()
    audit_type = audit_type_var.get()
    data_podpisu_umowy = data_podpisu_umowy_var.get().strip()
    data_podpisu_badania = data_podpisu_badania_var.get().strip()
    if not client_name or not year or not audit_type or not data_podpisu_umowy:
        messagebox.showerror("Error", "Proszę wypełnić wszystkie pola.")
        return
    # Use the spelling from the list even when the name was typed differently.
    client_name = next((name for name in client_names
                        if normalize_name(name) == normalize_name(client_name)), client_name)
    dzien_otw_bil = f"01.01.{year}"
    dzien_bil = f"31.12.{year}"
    if custom_dates_var.get():
        dzien_otw_bil = dzien_otw_bil_var.get().strip()
        dzien_bil = dzien_bil_var.get().strip()
    elif not re.fullmatch(r"\d{4}", year):
        messagebox.showerror("Error", "Nieprawidłowy rok badania.")
        return

    # Check every date now, before the signer selection, rather than at the very end.
    dates_to_check = [("Data podpisu umowy", data_podpisu_umowy)]
    if custom_dates_var.get():
        dates_to_check += [("Dzień otwarcia bilansu", dzien_otw_bil), ("Dzień bilansowy", dzien_bil)]
    if na_dzien_podpisu_var.get():
        dates_to_check.append(("Data podpisu SzB", data_podpisu_badania))
    invalid = [label for label, text in dates_to_check if parse_date(text) is None]
    if invalid:
        messagebox.showerror("Error", "Nieprawidłowy format daty (użyj formatu DD.MM.YYYY):\n" + "\n".join(invalid))
        return
    if parse_date(dzien_otw_bil) >= parse_date(dzien_bil):
        messagebox.showerror("Error", "Dzień otwarcia bilansu musi być wcześniejszy niż dzień bilansowy.")
        return

    save_settings(signature_color=signature_color_var.get(), audit_type=audit_type)
    client_entries = merged_clients if merged_clients else [client_name]
    open_signer_selection_dialog(client_name, client_entries, dzien_otw_bil, dzien_bil, audit_type,
                                 parse_date(data_podpisu_umowy), data_podpisu_badania)

# ------------------------- Consolidating Client Entries -------------------------
# Several entries of the list that are the same client (e.g. an old and a new name), treated as
# one client. The first one is the main entry: its name goes on the statement and in the file name.
merged_clients = []

def toggle_consolidate():
    if consolidate_var.get():
        open_consolidation_dialog()
    else:
        set_merged_clients([])

def set_merged_clients(names):
    merged_clients[:] = names
    if names:
        client_name_var.set(names[0])
        client_name_combobox.config(state=tk.DISABLED)  # untick "Połącz wpisy" to change the client
        merged_label.config(text=f"Połączone wpisy ({len(names)}): " + "; ".join(names))
        merged_label.grid()
    else:
        consolidate_var.set(0)
        client_name_combobox.config(state=tk.NORMAL)
        merged_label.grid_remove()

def open_consolidation_dialog():
    current = normalize_name(client_name_var.get())
    chosen = [name for name in client_names if normalize_name(name) == current]  # in the order picked
    shown = []

    dialog = tk.Toplevel(root)
    dialog.title("Połącz wpisy klienta")
    dialog.transient(root)
    ttk.Label(dialog, text="Zaznacz wpisy dotyczące tego samego klienta (kliknięcie zaznacza lub odznacza):"
              ).pack(anchor=tk.W, padx=10, pady=(10, 0))
    search_var = tk.StringVar()
    search_entry = ttk.Entry(dialog, textvariable=search_var)
    search_entry.pack(fill=tk.X, padx=10, pady=5)
    list_frame = ttk.Frame(dialog)
    list_frame.pack(fill=tk.BOTH, expand=True, padx=10)
    listbox = tk.Listbox(list_frame, selectmode=tk.MULTIPLE, exportselection=False,
                         width=80, height=15, font=('Calibri', 11))
    scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=listbox.yview)
    listbox.config(yscrollcommand=scrollbar.set)
    listbox.pack(side="left", fill=tk.BOTH, expand=True)
    scrollbar.pack(side="right", fill="y")
    chosen_label = ttk.Label(dialog, wraplength=600)
    chosen_label.pack(anchor=tk.W, padx=10, pady=5)

    def show_chosen():
        chosen_label.config(text=f"Wybrane ({len(chosen)}): " + "; ".join(chosen))

    def filter_list(*_):
        value = normalize_name(search_var.get())
        shown[:] = [name for name in client_names if value in normalize_name(name)]
        listbox.delete(0, tk.END)
        for i, name in enumerate(shown):
            listbox.insert(tk.END, name)
            if name in chosen:
                listbox.selection_set(i)
                listbox.see(i)

    def on_select(event):
        selected = set(listbox.curselection())
        for i, name in enumerate(shown):
            if i in selected and name not in chosen:
                chosen.append(name)
            elif i not in selected and name in chosen:
                chosen.remove(name)
        show_chosen()

    def on_ok():
        if len(chosen) < 2:
            messagebox.showerror("Error", "Wybierz co najmniej dwa wpisy klienta.", parent=dialog)
            return
        dialog.destroy()
        set_merged_clients(chosen)

    def on_cancel():
        dialog.destroy()
        set_merged_clients([])

    search_var.trace_add("write", filter_list)
    listbox.bind("<<ListboxSelect>>", on_select)
    dialog.protocol("WM_DELETE_WINDOW", on_cancel)
    btn_frame = ttk.Frame(dialog)
    btn_frame.pack(pady=10)
    ttk.Button(btn_frame, text="OK", command=on_ok).grid(row=0, column=0, padx=10)
    ttk.Button(btn_frame, text="Anuluj", command=on_cancel).grid(row=0, column=1, padx=10)
    filter_list()
    show_chosen()
    search_entry.focus_set()
    dialog.grab_set()

def toggle_custom_dates():
    state = tk.NORMAL if custom_dates_var.get() else tk.DISABLED
    dzien_otw_bil_box.config(state=state)
    dzien_bil_box.config(state=state)

def toggle_data_podpisu():
    state = tk.NORMAL if na_dzien_podpisu_var.get() else tk.DISABLED
    data_podpisu_badania_box.config(state=state)

def on_client_name_entry(event):
    value = normalize_name(client_name_var.get())
    if len(value) < 3:
        client_name_combobox['values'] = []
    else:
        data_list = [item for item in client_names if value in normalize_name(item)]
        client_name_combobox['values'] = data_list
        client_name_combobox.event_generate('<Down>')

# ------------------------- Main GUI Setup -------------------------
settings = load_settings()
root = tk.Tk()
root.title("Oświadczenia DAA")
root.minsize(500, 500)

# Apply ttk styles for a modern look
style = ttk.Style(root)
style.theme_use('clam')
style.configure('TLabel', font=('Calibri', 11), padding=5)
style.configure('TButton', font=('Calibri', 11), padding=5)
style.configure('TEntry', font=('Calibri', 11))
style.configure('TCombobox', font=('Calibri', 11))
style.configure('TLabelframe', font=('Calibri', 11, 'bold'), padding=10)
style.configure('TLabelframe.Label', font=('Calibri', 12, 'bold'))

main_frame = ttk.Frame(root, padding="15")
main_frame.grid(row=0, column=0, sticky=(tk.N, tk.S, tk.E, tk.W))
root.columnconfigure(0, weight=1)
root.rowconfigure(0, weight=1)
main_frame.columnconfigure(0, weight=1)
main_frame.columnconfigure(1, weight=2)

title_label = ttk.Label(main_frame, text="Generator Oświadczeń DAA", font=("Calibri", 16, "bold"))
title_label.grid(row=0, column=0, columnspan=2, pady=(0, 10))

grupa_kapitalowa_var = tk.IntVar()
grupa_kapitalowa_checkbutton = ttk.Checkbutton(main_frame, text="Grupa kapitałowa", variable=grupa_kapitalowa_var)
grupa_kapitalowa_checkbutton.grid(row=1, column=0, sticky=tk.W)

separator = ttk.Separator(main_frame, orient='horizontal')
separator.grid(row=2, column=0, columnspan=2, sticky='ew', pady=10)

client_frame = ttk.LabelFrame(main_frame, text="Dane oświadczenia")
client_frame.grid(row=3, column=0, columnspan=2, sticky=(tk.W, tk.E), pady=10)

ttk.Label(client_frame, text="Wybierz klienta:").grid(row=0, column=0, padx=5, pady=5, sticky=(tk.N, tk.W))
client_entry_frame = ttk.Frame(client_frame)
client_entry_frame.grid(row=0, column=1, padx=5, pady=5, sticky=(tk.W, tk.E))
client_entry_frame.columnconfigure(0, weight=1)
client_name_var = tk.StringVar()
client_name_combobox = ttk.Combobox(client_entry_frame, textvariable=client_name_var)
client_name_combobox['values'] = []
client_name_combobox.grid(row=0, column=0, sticky=(tk.W, tk.E))
client_name_combobox.bind('<KeyRelease>', on_client_name_entry)
consolidate_var = tk.IntVar()
consolidate_checkbutton = ttk.Checkbutton(client_entry_frame, text="Połącz wpisy", variable=consolidate_var,
                                          command=toggle_consolidate)
consolidate_checkbutton.grid(row=0, column=1, padx=(10, 0))
merged_label = ttk.Label(client_entry_frame, foreground="gray", wraplength=400)
merged_label.grid(row=1, column=0, columnspan=2, sticky=tk.W)
merged_label.grid_remove()  # shown only while entries are consolidated
client_frame.columnconfigure(1, weight=1)

ttk.Label(client_frame, text="Rok badania:").grid(row=1, column=0, padx=5, pady=5, sticky=tk.W)
year_var = tk.StringVar()
year_combobox = ttk.Combobox(client_frame, textvariable=year_var)
year_combobox['values'] = [str(y) for y in range(2021, datetime.now().year + 1)]
year_combobox.grid(row=1, column=1, padx=5, pady=5, sticky=(tk.W, tk.E))

custom_dates_var = tk.IntVar()
custom_dates_checkbutton = ttk.Checkbutton(
    client_frame, 
    text="Przesunięty rok obrotowy", 
    variable=custom_dates_var, 
    command=toggle_custom_dates
)
custom_dates_checkbutton.grid(row=2, column=0, columnspan=2, padx=5, pady=5, sticky=tk.W)

ttk.Label(client_frame, text="Dzień otwarcia bilansu:").grid(row=3, column=0, padx=5, pady=5, sticky=tk.W)
dzien_otw_bil_var = tk.StringVar(value="01.01.2021")
dzien_otw_bil_box = ttk.Entry(client_frame, textvariable=dzien_otw_bil_var, state=tk.DISABLED)
dzien_otw_bil_box.grid(row=3, column=1, padx=5, pady=5, sticky=(tk.W, tk.E))

ttk.Label(client_frame, text="Dzień bilansowy:").grid(row=4, column=0, padx=5, pady=5, sticky=tk.W)
dzien_bil_var = tk.StringVar(value="31.12.2021")
dzien_bil_box = ttk.Entry(client_frame, textvariable=dzien_bil_var, state=tk.DISABLED)
dzien_bil_box.grid(row=4, column=1, padx=5, pady=5, sticky=(tk.W, tk.E))

ttk.Label(client_frame, text="Data podpisu umowy:").grid(row=5, column=0, padx=5, pady=5, sticky=tk.W)
data_podpisu_umowy_var = tk.StringVar()
data_podpisu_umowy_box = ttk.Entry(client_frame, textvariable=data_podpisu_umowy_var)
data_podpisu_umowy_box.grid(row=5, column=1, padx=5, pady=5, sticky=(tk.W, tk.E))

ttk.Label(client_frame, text="Badanie/Przegląd:").grid(row=6, column=0, padx=5, pady=5, sticky=tk.W)
audit_type_var = tk.StringVar(value=settings.get("audit_type", ""))
audit_type_combobox = ttk.Combobox(client_frame, textvariable=audit_type_var)
audit_type_combobox['values'] = [
    "badaniem ",
    "badaniem jednostkowego ",
    "badaniem skonsolidowanego ",
    "badaniem jednostkowego i skonsolidowanego ",
    "badaniem śródrocznego skróconego jednostkowego ",
    "przeglądem śródrocznego jednostkowego i skonsolidowanego ",
    "przeglądem śródrocznego ",
    "przeglądem śródrocznego jednostkowego ",
    "przeglądem śródrocznego skonsolidowanego ",
    "przeglądem śródrocznego skróconego ",
    "przeglądem śródrocznego skróconego jednostkowego ",
    "przeglądem śródrocznego skróconego skonsolidowanego "
]
audit_type_combobox.grid(row=6, column=1, padx=5, pady=5, sticky=(tk.W, tk.E))

ttk.Label(client_frame, text="Kolor podpisu:").grid(row=7, column=0, padx=5, pady=5, sticky=tk.W)
signature_color_var = tk.StringVar(value=settings.get("signature_color", "czarny"))
signature_color_combobox = ttk.Combobox(client_frame, textvariable=signature_color_var)
signature_color_combobox['values'] = ["czarny", "niebieski"]
signature_color_combobox.grid(row=7, column=1, padx=5, pady=5, sticky=(tk.W, tk.E))

na_dzien_podpisu_var = tk.IntVar()
na_dzien_podpisu_checkbutton = ttk.Checkbutton(
    client_frame, 
    text="Na dzień podpisu SzB", 
    variable=na_dzien_podpisu_var, 
    command=toggle_data_podpisu
)
na_dzien_podpisu_checkbutton.grid(row=8, column=0, columnspan=2, padx=5, pady=5, sticky=tk.W)

ttk.Label(client_frame, text="Data podpisu SzB:").grid(row=9, column=0, padx=5, pady=5, sticky=tk.W)
data_podpisu_badania_var = tk.StringVar(value=datetime.today().strftime("%d.%m.%Y"))
data_podpisu_badania_box = ttk.Entry(client_frame, textvariable=data_podpisu_badania_var, state=tk.DISABLED)
data_podpisu_badania_box.grid(row=9, column=1, padx=5, pady=5, sticky=(tk.W, tk.E))

separator2 = ttk.Separator(main_frame, orient='horizontal')
separator2.grid(row=4, column=0, columnspan=2, sticky='ew', pady=10)

button_frame = ttk.Frame(main_frame)
button_frame.grid(row=5, column=0, columnspan=2, pady=10)
submit_button = ttk.Button(button_frame, text="Dalej", command=submit_form)
submit_button.grid(row=0, column=0, padx=10)
refresh_button = ttk.Button(button_frame, text="Odśwież dane", command=refresh_data)
refresh_button.grid(row=0, column=1, padx=10)
exit_button = ttk.Button(button_frame, text="Wyjdź", command=root.quit)
exit_button.grid(row=0, column=2, padx=10)
button_frame.columnconfigure((0, 1, 2), weight=1)

status_var = tk.StringVar()
ttk.Label(main_frame, textvariable=status_var, foreground="gray").grid(row=6, column=0, columnspan=2)

# Show the local copy right away, then fetch what changed in SharePoint since the last run.
load_data()
show_data_status()
root.after(100, refresh_data, False)  # once the window is shown, so the sign-in window can open over it
root.mainloop()
